import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

BOOK = Path.home() / "Projects/stelios/dnd/characters/laynlindr-freth/backstory/book"
OUT = BOOK / "Soulknife.docx"
BODY = "Garamond"
ORNAMENT = "\u2020  \u2020"  # paired daggers


def base(doc):
    s = doc.styles["Normal"]
    s.font.name = BODY
    s.font.size = Pt(11.5)
    pf = s.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Inches(0.3)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    for sec in doc.sections:
        sec.page_width = Inches(6)
        sec.page_height = Inches(9)
        sec.top_margin = sec.bottom_margin = Inches(0.75)
        sec.left_margin = sec.right_margin = Inches(0.75)


FN_ID = 2


def add_footnote_part(doc, note_text):
    W = nsdecls("w")
    xml = (
        f'<w:footnotes {W}>'
        f'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        f'<w:footnote w:id="{FN_ID}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        f'<w:r><w:rPr><w:vertAlign w:val="superscript"/><w:sz w:val="18"/></w:rPr><w:footnoteRef/></w:r>'
        f'<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:t xml:space="preserve"> {note_text}</w:t></w:r></w:p></w:footnote>'
        f'</w:footnotes>'
    )
    part = Part(PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                xml.encode(), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)


def footnote_ref(p):
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
        f'<w:footnoteReference w:id="{FN_ID}"/></w:r>'))


def curly(text):
    return re.sub(r'(\s|^)"', lambda m: m.group(1) + "\u201c", text).replace('"', "\u201d")


def runs(p, text):
    # inline **bold** / *italic* / [^1] footnote ref
    for i, chunk in enumerate(re.split(r"\[\^1\]", curly(text))):
        if i:
            footnote_ref(p)
        for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", chunk):
            if not part:
                continue
            if part.startswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("*"):
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)


def chapter(doc, num, title, paras):
    doc.add_page_break()
    for txt, size, before in [
        *([("Chapter " + num, 14, 60)] if num else []),
        (title, 22 if num else 18, 60 if num else 80),
        (ORNAMENT, 12, 14),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(txt)
        r.font.name = BODY
        r.font.size = Pt(size)
        if txt != ORNAMENT:
            r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    doc.add_paragraph().paragraph_format.first_line_indent = None
    after_break = False
    for para in paras:
        if para.startswith("> "):  # blockquote: in-world document, indented italic
            q = doc.add_paragraph()
            q.paragraph_format.first_line_indent = None
            q.paragraph_format.left_indent = Inches(0.35)
            q.paragraph_format.space_before = Pt(6)
            q.paragraph_format.space_after = Pt(6)
            r = q.add_run(curly(para[2:]))
            r.italic = True
            after_break = False
            continue
        if para == "---":  # scene break -> blank space
            sp = doc.add_paragraph()
            sp.paragraph_format.first_line_indent = None
            sp.paragraph_format.space_before = Pt(12)
            after_break = True
            continue
        p = doc.add_paragraph()
        if after_break:
            p.paragraph_format.first_line_indent = None
            after_break = False
        runs(p, para)


def centered(doc, lines):
    doc.add_paragraph()
    for txt, size, before in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(txt)
        r.font.size = Pt(size)
        if txt.isupper():
            r.bold = True


doc = Document()
base(doc)

title_lines = [l.lstrip("# ").strip() for l in (BOOK / "chapter-00-title.md").read_text().splitlines() if l.strip()]
centered(doc, [(title_lines[0].upper(), 28, 140), (title_lines[1].upper(), 13, 6), (ORNAMENT, 14, 20)])

# map page after the title
if (BOOK / "menzoberranzan-map.png").exists():
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(str(BOOK / "menzoberranzan-map.png"), width=Inches(4.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = None
    r = cap.add_run("Menzoberranzan — map \u00a9 Wizards of the Coast, \u201cMenzoberranzan: City of Intrigue\u201d")
    r.font.size = Pt(8)
    r.italic = True

chapters = []
for f in sorted(BOOK.glob("chapter-*.md")):
    if "title" in f.name:
        continue
    text = f.read_text()
    m = re.match(r"#+\s+(?:Chapter\s+(\d+))?\s*[—–-]?\s*(.+)", text)
    num, name = (m.group(1), m.group(2).strip()) if m and m.group(1) else (None, m.group(2).strip() if m else f.stem)
    paras = [ln.strip() for ln in text.splitlines()[1:]
             if ln.strip() and not ln.strip().startswith("[^1]:")]
    if "readers-guide" in f.name:  # glossary goes at the end
        glossary = (num, name, paras)
        continue
    chapters.append((num, name, paras))

for num, name, paras in chapters:
    chapter(doc, num, name, paras)

# glossary as back matter, no first-line indent
num, name, paras = glossary
chapter(doc, num, name, paras)
for p in doc.paragraphs[-len(glossary[2]):]:
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(6)

if any("[^1]" in p for _, _, ps in chapters + [glossary] for p in ps):
    add_footnote_part(doc, "\u201cAstux\u201d is drowish for \u201cextinguish.\u201d")

doc.save(OUT)
print(OUT)